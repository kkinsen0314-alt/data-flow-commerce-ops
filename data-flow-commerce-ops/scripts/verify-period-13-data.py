from __future__ import annotations

import csv
import json
import re
import sys
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURE_ROOT = ROOT / "data" / "fixtures"
PERIOD_ROOT = FIXTURE_ROOT / "period-13"
WORKBOOK_PATH = ROOT / "outputs" / "period-13-data" / "period-13-commerce-data.xlsx"
ARTIFACT_PATH = ROOT / "artifacts" / "period-13-data-validation-v1.json"
DATETIME_FORMAT = "%Y-%m-%d %H:%M:%S"

FILES = {
    "short_video": {
        "path": PERIOD_ROOT / "short-video-period-13.csv",
        "sheet": "短视频",
        "headers": [
            "content_id_hash", "account_id_hash", "click_id_hash",
            "published_at", "impressions", "plays", "completions",
            "interactions", "clicks",
        ],
        "rows": 18,
    },
    "live_session": {
        "path": PERIOD_ROOT / "live-session-period-13.csv",
        "sheet": "直播",
        "headers": [
            "用户ID", "直播场次", "直播标题", "是否到课", "是否完课",
            "是否访问课程商品", "是否领取优惠券", "是否发起支付",
            "是否购买课程", "是否观看回放", "直播观看时长", "邀请人",
            "线索渠道", "是否新用户", "预约时间",
        ],
        "rows": 24,
    },
    "channel_lead": {
        "path": PERIOD_ROOT / "channel-leads-period-13.csv",
        "sheet": "渠道线索",
        "headers": [
            "lead_id_hash", "click_id_hash", "sales_owner_id_hash", "channel",
            "lead_source", "created_at", "lead_stage",
        ],
        "rows": 18,
    },
    "sales_followup": {
        "path": PERIOD_ROOT / "sales-followup-period-13.csv",
        "sheet": "销售跟进",
        "headers": [
            "lead_id_hash", "sales_owner_id_hash", "assigned_at",
            "first_followup_at", "followup_count", "followup_status",
        ],
        "rows": 17,
    },
    "order": {
        "path": PERIOD_ROOT / "orders-period-13.csv",
        "sheet": "订单",
        "headers": [
            "order_id_hash", "lead_id_hash", "ordered_at", "paid_at",
            "paid_amount", "order_status",
        ],
        "rows": 12,
    },
}


def read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def parse_datetime(value: str) -> datetime:
    return datetime.strptime(value, DATETIME_FORMAT)


def canonical_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(DATETIME_FORMAT)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


checks: list[dict[str, Any]] = []


def record(name: str, passed: bool, details: Any) -> None:
    checks.append({"name": name, "status": "pass" if passed else "fail", "details": details})


loaded: dict[str, list[dict[str, str]]] = {}
for dataset_type, config in FILES.items():
    path = config["path"]
    record(f"{dataset_type}.file_exists", path.is_file(), str(path))
    if not path.is_file():
        loaded[dataset_type] = []
        continue
    headers, rows = read_csv_rows(path)
    loaded[dataset_type] = rows
    record(f"{dataset_type}.headers", headers == config["headers"], headers)
    record(f"{dataset_type}.row_count", len(rows) == config["rows"], len(rows))
    signatures = [tuple(row.get(header, "") for header in headers) for row in rows]
    record(f"{dataset_type}.no_duplicate_rows", len(signatures) == len(set(signatures)), len(signatures) - len(set(signatures)))

video = loaded["short_video"]
live = loaded["live_session"]
leads = loaded["channel_lead"]
followups = loaded["sales_followup"]
orders = loaded["order"]

video_ids = [row["content_id_hash"] for row in video]
video_clicks = {row["click_id_hash"] for row in video}
lead_ids = {row["lead_id_hash"] for row in leads}
lead_clicks = {row["click_id_hash"] for row in leads}
followup_ids = {row["lead_id_hash"] for row in followups}
order_ids = [row["order_id_hash"] for row in orders]
order_leads = {row["lead_id_hash"] for row in orders}
live_leads = {row["用户ID"] for row in live}

record("short_video.unique_primary_ids", len(video_ids) == len(set(video_ids)), len(video_ids))
record("orders.unique_primary_ids", len(order_ids) == len(set(order_ids)), len(order_ids))
record("click_join.exact_coverage", video_clicks == lead_clicks and len(video_clicks) == 18, len(video_clicks & lead_clicks))
record("live_to_lead.full_coverage", live_leads == lead_ids, len(live_leads & lead_ids))
record("followup_to_lead.valid_subset", followup_ids <= lead_ids, sorted(followup_ids - lead_ids))
record("order_to_lead.valid_subset", order_leads <= lead_ids, sorted(order_leads - lead_ids))
record("followup.intentional_gap", lead_ids - followup_ids == {"p13_lead_017"}, sorted(lead_ids - followup_ids))

lead_owner = {row["lead_id_hash"]: row["sales_owner_id_hash"] for row in leads}
owner_mismatches = [row["lead_id_hash"] for row in followups if lead_owner.get(row["lead_id_hash"]) != row["sales_owner_id_hash"]]
record("followup.owner_consistency", not owner_mismatches, owner_mismatches)

funnel_errors = []
for row in video:
    impressions, plays, completions, interactions, clicks = (
        int(row["impressions"]), int(row["plays"]), int(row["completions"]),
        int(row["interactions"]), int(row["clicks"]),
    )
    if not (impressions >= plays >= completions >= 0 and interactions >= clicks >= 0):
        funnel_errors.append(row["content_id_hash"])
record("short_video.funnel_monotonic", not funnel_errors, funnel_errors)

boolean_columns = [
    "是否到课", "是否完课", "是否访问课程商品", "是否领取优惠券",
    "是否发起支付", "是否购买课程", "是否观看回放",
]
invalid_booleans = [
    f"{row['用户ID']}:{column}"
    for row in live
    for column in boolean_columns
    if row[column] not in {"是", "否"}
]
record("live.boolean_values", not invalid_booleans, invalid_booleans)
record("live.sessions", {row["直播场次"] for row in live} == {"1301", "1302"}, sorted({row["直播场次"] for row in live}))
record("live.period_title", all("第13期课程" in row["直播标题"] for row in live), sorted({row["直播标题"] for row in live}))

duration_errors = []
for row in live:
    try:
        hours, minutes, seconds = [int(part) for part in row["直播观看时长"].split(":")]
        valid = hours >= 0 and 0 <= minutes < 60 and 0 <= seconds < 60
    except (TypeError, ValueError):
        valid = False
    if not valid:
        duration_errors.append(row["用户ID"])
record("live.duration_format", not duration_errors, duration_errors)

all_timestamps: list[datetime] = []
timestamp_errors: list[str] = []
timestamp_fields = {
    "short_video": ["published_at"],
    "live_session": ["预约时间"],
    "channel_lead": ["created_at"],
    "sales_followup": ["assigned_at", "first_followup_at"],
    "order": ["ordered_at", "paid_at"],
}
for dataset_type, fields in timestamp_fields.items():
    for index, row in enumerate(loaded[dataset_type], start=1):
        for field in fields:
            value = row[field]
            if not value:
                continue
            try:
                all_timestamps.append(parse_datetime(value))
            except ValueError:
                timestamp_errors.append(f"{dataset_type}:{index}:{field}")
record("timestamps.parseable", not timestamp_errors, timestamp_errors)
timestamp_min = min(all_timestamps) if all_timestamps else None
timestamp_max = max(all_timestamps) if all_timestamps else None
record(
    "timestamps.period_window",
    timestamp_min == datetime(2026, 8, 31, 9, 0, 0) and timestamp_max == datetime(2026, 9, 6, 22, 30, 0),
    {"min": timestamp_min, "max": timestamp_max},
)

lead_created = {row["lead_id_hash"]: parse_datetime(row["created_at"]) for row in leads}
followup_sequence_errors = []
within_24h = 0
for row in followups:
    lead_id = row["lead_id_hash"]
    assigned = parse_datetime(row["assigned_at"])
    first = parse_datetime(row["first_followup_at"])
    if not (lead_created[lead_id] <= assigned <= first):
        followup_sequence_errors.append(lead_id)
    if first - assigned <= timedelta(hours=24):
        within_24h += 1
record("followup.time_sequence", not followup_sequence_errors, followup_sequence_errors)
record("followup.within_24h", within_24h == 13, {"numerator": within_24h, "denominator": len(leads), "ratio": within_24h / len(leads)})

order_sequence_errors = []
paid_at_errors = []
for row in orders:
    ordered_at = parse_datetime(row["ordered_at"])
    paid_at = parse_datetime(row["paid_at"]) if row["paid_at"] else None
    if ordered_at < lead_created[row["lead_id_hash"]] or (paid_at and paid_at < ordered_at):
        order_sequence_errors.append(row["order_id_hash"])
    if (row["order_status"] in {"paid", "completed"}) != bool(paid_at):
        paid_at_errors.append(row["order_id_hash"])
record("orders.time_sequence", not order_sequence_errors, order_sequence_errors)
record("orders.paid_at_consistency", not paid_at_errors, paid_at_errors)

paid_orders = [row for row in orders if row["order_status"] in {"paid", "completed"}]
paid_gmv = sum(Decimal(row["paid_amount"]) for row in paid_orders)
record("metrics.paid_order_count", len(paid_orders) == 9, len(paid_orders))
record("metrics.paid_gmv", paid_gmv == Decimal("10591"), paid_gmv)
record("metrics.order_lead_coverage", len(order_leads) == 12 and len(order_leads) / len(leads) == 2 / 3, len(order_leads) / len(leads))
record("metrics.paid_lead_conversion", len(paid_orders) / len(leads) == 0.5, len(paid_orders) / len(leads))

all_text = "\n".join(value for rows in loaded.values() for row in rows for value in row.values())
privacy_hits = {
    "mobile": re.findall(r"(?<!\d)1[3-9]\d{9}(?!\d)", all_text),
    "email": re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", all_text),
}
identifier_values = [
    value
    for rows in loaded.values()
    for row in rows
    for key, value in row.items()
    if "ID" in key or key.endswith("_hash")
]
record("privacy.no_phone_or_email", not privacy_hits["mobile"] and not privacy_hits["email"], privacy_hits)
record("privacy.synthetic_id_namespace", all(value.startswith("p13_") for value in identifier_values), len(identifier_values))

record("workbook.file_exists", WORKBOOK_PATH.is_file(), str(WORKBOOK_PATH))
workbook_differences: list[str] = []
formula_cells: list[str] = []
sheet_names: list[str] = []
if WORKBOOK_PATH.is_file():
    workbook = load_workbook(WORKBOOK_PATH, data_only=False, read_only=True)
    sheet_names = workbook.sheetnames
    for dataset_type, config in FILES.items():
        sheet = workbook[config["sheet"]]
        csv_headers, csv_rows = read_csv_rows(config["path"])
        worksheet_rows = list(sheet.iter_rows(values_only=True))
        worksheet_text = [[canonical_cell(value) for value in row] for row in worksheet_rows]
        csv_text = [csv_headers] + [[row[header] for header in csv_headers] for row in csv_rows]
        if worksheet_text != csv_text:
            workbook_differences.append(dataset_type)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(f"{config['sheet']}!{cell.coordinate}")
record("workbook.sheet_set", sheet_names == [config["sheet"] for config in FILES.values()], sheet_names)
record("workbook.csv_exact_parity", not workbook_differences, workbook_differences)
record("workbook.no_formulas", not formula_cells, formula_cells)

production_manifests: dict[str, Any] = {}
native_normalized: list[dict[str, Any]] = []
production_error: str | None = None
try:
    sys.path.insert(0, str(ROOT))
    from commerce_ops.datasets import DatasetStore
    from commerce_ops.native_models import NativeRunRequest
    from commerce_ops.native_runtime import NativeRuntime
    from commerce_ops.tool_models import DataReference

    store = DatasetStore(FIXTURE_ROOT)
    for index, (dataset_type, config) in enumerate(FILES.items(), start=1):
        reference = DataReference(
            dataset_id=f"ds_p13_{index:02d}",
            dataset_type=dataset_type,
            file_path=f"period-13/{config['path'].name}",
            synthetic=True,
        )
        manifest = store.load("wf_period13_validation", reference).manifest
        production_manifests[dataset_type] = {
            "quality_status": manifest.data_quality.status,
            "row_count": manifest.data_quality.row_count,
            "warnings": manifest.data_quality.warnings,
            "missing_required_fields": manifest.data_quality.missing_required_fields,
        }

    request = NativeRunRequest(
        idempotency_key="period13-validation-only",
        requested_domains=["content_growth", "live_conversion", "attribution_leads"],
        datasets=[
            {"dataset_type": dataset_type, "file_path": f"period-13/{config['path'].name}"}
            for dataset_type, config in FILES.items()
        ],
        objective="仅执行本地第13期数据路径与域覆盖验证，不启动原生运行。",
        include_strategy=True,
        authorized_model_execution=True,
    )
    runtime = NativeRuntime(
        project_root=ROOT,
        data_root=FIXTURE_ROOT,
        workspace_jid=None,
        client=None,
        store=None,
        audit_reader=None,
    )
    native_normalized = [item.model_dump(mode="json") for item in runtime._normalize_datasets(request)]
except Exception as exc:
    production_error = f"{type(exc).__name__}: {exc}"

record(
    "production_loader.all_manifests_pass",
    production_error is None
    and len(production_manifests) == 5
    and all(item["quality_status"] == "pass" for item in production_manifests.values()),
    production_manifests if production_error is None else production_error,
)
record(
    "native_preflight.paths_and_domains",
    production_error is None and len(native_normalized) == 5,
    native_normalized if production_error is None else production_error,
)

failed = [item for item in checks if item["status"] != "pass"]
artifact = {
    "schema_version": "1.0",
    "artifact_type": "period_13_synthetic_data_validation",
    "status": "pass" if not failed else "fail",
    "synthetic": True,
    "period": {"label": "第13期", "start": "2026-08-31 09:00:00", "end": "2026-09-06 22:30:00"},
    "execution_guards": {
        "host_started": False,
        "api_request_sent": False,
        "model_called": False,
        "native_run_created": False,
    },
    "row_counts": {dataset_type: len(rows) for dataset_type, rows in loaded.items()},
    "business_metrics": {
        "paid_order_count": len(paid_orders),
        "paid_gmv": str(paid_gmv),
        "order_lead_coverage": len(order_leads) / len(leads),
        "paid_lead_conversion": len(paid_orders) / len(leads),
        "followup_within_24h": {"numerator": within_24h, "denominator": len(leads), "ratio": within_24h / len(leads)},
        "intentional_missing_followup_lead_ids": sorted(lead_ids - followup_ids),
    },
    "production_manifests": production_manifests,
    "native_normalized_datasets": native_normalized,
    "checks_total": len(checks),
    "checks_passed": len(checks) - len(failed),
    "checks_failed": len(failed),
    "checks": checks,
}
ARTIFACT_PATH.parent.mkdir(parents=True, exist_ok=True)
ARTIFACT_PATH.write_text(json.dumps(artifact, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
print(json.dumps({
    "status": artifact["status"],
    "checks_total": artifact["checks_total"],
    "checks_passed": artifact["checks_passed"],
    "checks_failed": artifact["checks_failed"],
    "artifact": str(ARTIFACT_PATH),
}, ensure_ascii=False))
raise SystemExit(1 if failed else 0)
